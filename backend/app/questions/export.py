"""Question set exports for interviewers: Markdown, DOCX (python-docx) and PDF (reportlab).

They contain ideal answers and red flags, so every export is marked as a confidential interviewer copy.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Any
from xml.sax.saxutils import escape

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

CONFIDENTIAL = "Confidential: interviewer copy. Contains model answers and red flags; do not share with candidates."


@dataclass(frozen=True, slots=True)
class ExportSet:
    title: str
    subtitle: str
    items: list[dict[str, Any]]
    competency_names: dict[str, str]
    citation_titles: dict[str, str]


def _sections(item: dict[str, Any], doc: ExportSet) -> list[tuple[str, list[str] | str]]:
    citations = [doc.citation_titles.get(c, c) for c in item.get("citations", [])]
    return [
        ("Why ask", item.get("rationale", "")),
        ("Strong answers mention", item.get("expected_signals", [])),
        ("Red flags", item.get("red_flags", [])),
        ("Follow-ups", item.get("follow_ups", [])),
        ("Model answer", item.get("ideal_answer", "")),
        ("Sources", citations),
    ]


def _meta(item: dict[str, Any], doc: ExportSet) -> str:
    competency = doc.competency_names.get(item.get("competency", ""), item.get("competency", ""))
    return (
        f"{competency} · {str(item.get('category', '')).replace('_', ' ')} · difficulty {item.get('difficulty', '?')}/5"
    )


def to_markdown(doc: ExportSet) -> str:
    lines = [f"# {doc.title}", "", f"_{doc.subtitle}_", "", f"> {CONFIDENTIAL}", ""]
    for number, item in enumerate(doc.items, start=1):
        lines += [f"## {number}. {item['text']}", "", f"*{_meta(item, doc)}*", ""]
        for label, value in _sections(item, doc):
            if not value:
                continue
            if isinstance(value, list):
                lines += [f"**{label}:**", *[f"- {v}" for v in value], ""]
            else:
                lines += [f"**{label}:** {value}", ""]
    return "\n".join(lines).rstrip() + "\n"


def to_docx(doc: ExportSet) -> bytes:
    document = DocxDocument()
    document.add_heading(doc.title, level=0)
    document.add_paragraph().add_run(doc.subtitle).italic = True
    document.add_paragraph(CONFIDENTIAL)
    for number, item in enumerate(doc.items, start=1):
        document.add_heading(f"{number}. {item['text']}", level=2)
        document.add_paragraph().add_run(_meta(item, doc)).italic = True
        for label, value in _sections(item, doc):
            if not value:
                continue
            if isinstance(value, list):
                document.add_paragraph(label).runs[0].bold = True
                for entry in value:
                    document.add_paragraph(str(entry), style="List Bullet")
            else:
                paragraph = document.add_paragraph()
                paragraph.add_run(f"{label}: ").bold = True
                paragraph.add_run(str(value))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def to_pdf(doc: ExportSet) -> bytes:
    styles = getSampleStyleSheet()
    body, small = styles["BodyText"], styles["Italic"]
    story: list[Any] = [
        Paragraph(escape(doc.title), styles["Title"]),
        Paragraph(escape(doc.subtitle), small),
        Spacer(1, 4 * mm),
        Paragraph(f"<b>{escape(CONFIDENTIAL)}</b>", body),
        Spacer(1, 6 * mm),
    ]
    for number, item in enumerate(doc.items, start=1):
        story += [
            Paragraph(escape(f"{number}. {item['text']}"), styles["Heading3"]),
            Paragraph(escape(_meta(item, doc)), small),
        ]
        for label, value in _sections(item, doc):
            if not value:
                continue
            if isinstance(value, list):
                story.append(Paragraph(f"<b>{escape(label)}</b>", body))
                story.append(
                    ListFlowable(
                        [ListItem(Paragraph(escape(str(v)), body), leftIndent=10) for v in value], bulletType="bullet"
                    )
                )
            else:
                story.append(Paragraph(f"<b>{escape(label)}:</b> {escape(str(value))}", body))
        story.append(Spacer(1, 5 * mm))
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, title=doc.title, leftMargin=18 * mm, rightMargin=18 * mm).build(story)
    return buffer.getvalue()
