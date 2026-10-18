"""Document parsing to Markdown with page mapping, tables and risk signals.

Backends:
  basic   pdfplumber (PDF), python-docx, python-pptx, openpyxl, BeautifulSoup, plain text; RapidOCR for scanned
          pages and images when the `ocr` extra is installed. Fully offline.
  docling IBM Docling (better layout/tables/figures) when installed with its models; used by `auto` if present.
"""

from __future__ import annotations

import io
import re
import statistics
from dataclasses import dataclass, field
from typing import Any

import filetype

from app.core.logging import get_logger
from app.rag.risk import injection_matches

log = get_logger("app.rag.parse")

_HEADING_RATIO = 1.25
_MAX_HEADING_CHARS = 120
_MIN_VISIBLE_SIZE = 2.0
_WHITE = 0.95
_CMYK_COMPONENTS = 4
_XLSX_MAX_ROWS = 500


class ParseError(Exception):
    """The document cannot be parsed (unsupported, corrupt, encrypted, or needs an unavailable component)."""


@dataclass
class ParsedPage:
    number: int
    markdown: str


@dataclass
class ParsedDocument:
    pages: list[ParsedPage]
    title: str | None = None
    meta: dict[str, Any] = field(default_factory=dict)
    hidden_text: list[str] = field(default_factory=list)
    ocr_pages: list[int] = field(default_factory=list)

    @property
    def markdown(self) -> str:
        return "\n\n".join(p.markdown for p in self.pages if p.markdown.strip())

    def risk_flags(self) -> dict[str, Any]:
        flags: dict[str, Any] = {}
        if self.hidden_text:
            flags["hidden_text"] = self.hidden_text[:20]
        matches = injection_matches(self.markdown + "\n" + "\n".join(self.hidden_text))
        if matches:
            flags["injection_patterns"] = matches
        if self.ocr_pages:
            flags["ocr_pages"] = self.ocr_pages
        return flags


# --- type detection ----------------------------------------------------------------------------------

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_EXTENSIONS = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "html": "text/html",
    "htm": "text/html",
    "md": "text/markdown",
    "markdown": "text/markdown",
    "txt": "text/plain",
    "csv": "text/csv",
    "c": "text/x-c",
    "h": "text/x-c",
    "cpp": "text/x-c++",
    "hpp": "text/x-c++",
    "py": "text/x-python",
    "rs": "text/x-rust",
    "json": "application/json",
    "yaml": "application/yaml",
    "yml": "application/yaml",
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "webp": "image/webp",
}
SUPPORTED_MIME = frozenset(_EXTENSIONS.values())


def detect_mime(data: bytes, filename: str | None) -> str:
    """Trust content signatures over the filename; fall back to the extension for text formats."""
    kind = filetype.guess(data[:8192])
    extension = (filename or "").rsplit(".", 1)[-1].lower() if filename and "." in filename else ""
    if kind is not None:
        if kind.mime == "application/zip" and extension in {"docx", "pptx", "xlsx"}:
            return _EXTENSIONS[extension]
        return str(kind.mime)
    if extension in _EXTENSIONS and _EXTENSIONS[extension].startswith(
        ("text/", "application/json", "application/yaml")
    ):
        return _EXTENSIONS[extension]
    sample = data[:65536]
    if b"\x00" in sample:  # binaries (firmware images, executables) are not text even if they decode
        return "application/octet-stream"
    try:
        sample.decode("utf-8")
    except UnicodeDecodeError:
        return "application/octet-stream"
    return "text/plain"


# --- helpers -------------------------------------------------------------------------------------------


def _table_markdown(rows: list[list[Any]]) -> str:
    cleaned = [
        [re.sub(r"\s+", " ", str(cell or "")).strip().replace("|", "\\|") for cell in row] for row in rows if row
    ]
    cleaned = [row for row in cleaned if any(row)]
    if not cleaned:
        return ""
    width = max(len(r) for r in cleaned)
    cleaned = [r + [""] * (width - len(r)) for r in cleaned]
    lines = ["| " + " | ".join(cleaned[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in cleaned[1:]]
    return "\n".join(lines)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def ocr_available() -> bool:
    try:
        import rapidocr_onnxruntime  # noqa: F401, PLC0415
    except ImportError:
        return False
    return True


def _ocr_image(image: Any) -> str:
    from rapidocr_onnxruntime import RapidOCR  # noqa: PLC0415 - optional extra

    engine = _ocr_engine_cache.get("engine")
    if engine is None:
        engine = RapidOCR()
        _ocr_engine_cache["engine"] = engine
    import numpy as np  # noqa: PLC0415

    result, _ = engine(np.array(image.convert("RGB")))
    if not result:
        return ""
    # Sort boxes top-to-bottom, left-to-right and join lines.
    boxes = sorted(result, key=lambda r: (round(r[0][0][1] / 12), r[0][0][0]))
    return "\n".join(str(r[1]) for r in boxes)


_ocr_engine_cache: dict[str, Any] = {}


# --- PDF -------------------------------------------------------------------------------------------------


def _is_white(color: Any) -> bool:
    if color is None:
        return False
    values = [float(c) for c in (color if isinstance(color, list | tuple) else [color]) if isinstance(c, int | float)]
    if not values:
        return False
    if len(values) == _CMYK_COMPONENTS:  # CMYK: white is no ink at all
        return all(v <= 1 - _WHITE for v in values)
    return all(v >= _WHITE for v in values)


def _hidden(char: dict[str, Any], width: float, height: float) -> bool:
    if float(char.get("size", 10)) < _MIN_VISIBLE_SIZE:
        return True
    if _is_white(char.get("non_stroking_color")):
        return True
    return bool(char["x1"] < 0 or char["x0"] > width or char["bottom"] < 0 or char["top"] > height)


def _pdf_lines(visible: Any, body_size: float) -> list[str]:
    lines: list[str] = []
    for line in visible.extract_text_lines(return_chars=True):
        sizes = [float(c["size"]) for c in line.get("chars", [])]
        text_line = line["text"].strip()
        if not text_line:
            continue
        is_heading = sizes and statistics.mean(sizes) >= body_size * _HEADING_RATIO
        if is_heading and len(text_line) <= _MAX_HEADING_CHARS:
            lines.append(f"\n## {text_line}\n")
        else:
            lines.append(text_line)
    return lines


def _pdf_page(page: Any, number: int, body_size: float) -> tuple[ParsedPage, str | None, bool]:
    """Returns (page, hidden text found on it, whether OCR was needed)."""
    width, height = float(page.width), float(page.height)
    hidden_chars = [c for c in page.chars if _hidden(c, width, height)]
    hidden_text = "".join(c["text"] for c in hidden_chars).strip()[:500] or None
    tables = page.find_tables()
    boxes = [t.bbox for t in tables]

    def keep(obj: dict[str, Any]) -> bool:
        if obj.get("object_type") != "char":
            return True
        if _hidden(obj, width, height):
            return False
        return not any(
            b[0] <= obj["x0"] and obj["x1"] <= b[2] and b[1] <= obj["top"] and obj["bottom"] <= b[3] for b in boxes
        )

    body = "\n".join(_pdf_lines(page.filter(keep), body_size)).strip()
    for table in tables:
        rendered = _table_markdown(table.extract())
        if rendered:
            body += "\n\n" + rendered
    needs_ocr = not body and not page.chars
    if needs_ocr and ocr_available():
        body = _ocr_image(page.to_image(resolution=200).original)
    return ParsedPage(number, body), hidden_text, needs_ocr


def _parse_pdf(data: bytes, max_pages: int) -> ParsedDocument:
    import pdfplumber  # noqa: PLC0415 - keeps import cost off the API's startup path

    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"Not a readable PDF ({type(exc).__name__}). It may be corrupt or encrypted.") from exc
    pages: list[ParsedPage] = []
    hidden: list[str] = []
    ocr_pages: list[int] = []
    with pdf:
        if len(pdf.pages) > max_pages:
            raise ParseError(f"The PDF has {len(pdf.pages)} pages; the limit is {max_pages}.")
        all_sizes = [
            float(c["size"])
            for page in pdf.pages[:20]
            for c in page.chars
            if float(c.get("size", 0)) >= _MIN_VISIBLE_SIZE
        ]
        body_size = statistics.median(all_sizes) if all_sizes else 10.0
        for number, page in enumerate(pdf.pages, start=1):
            parsed_page, hidden_text, needs_ocr = _pdf_page(page, number, body_size)
            if hidden_text:
                hidden.append(hidden_text)
            if needs_ocr:
                ocr_pages.append(number)
            pages.append(parsed_page)
        title = (pdf.metadata or {}).get("Title") or None
    if not any(p.markdown.strip() for p in pages):
        detail = "" if ocr_available() else " Install the OCR extra (or use a vision model) to read scanned documents."
        raise ParseError("No text could be extracted from this PDF." + detail)
    return ParsedDocument(pages=pages, title=title, hidden_text=hidden, ocr_pages=ocr_pages, meta={"pages": len(pages)})


# --- Office, web, text, images ---------------------------------------------------------------------------------


def _parse_docx(data: bytes) -> ParsedDocument:
    import docx  # noqa: PLC0415

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"Not a readable Word document ({type(exc).__name__}).") from exc
    parts: list[str] = []
    hidden: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        runs_hidden = [
            r.text
            for r in paragraph.runs
            if r.font.hidden or (r.font.color and r.font.color.rgb and str(r.font.color.rgb).upper() == "FFFFFF")
        ]
        if runs_hidden:
            hidden.append(" ".join(runs_hidden)[:500])
            text = "".join(r.text for r in paragraph.runs if r.text not in runs_hidden).strip()
            if not text:
                continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style.lower().startswith("heading") or style.lower() == "title":
            level = re.findall(r"\d+", style)
            parts.append("#" * min(6, (int(level[0]) if level else 1) + 1) + " " + text)
        elif style.lower().startswith("list"):
            parts.append(f"- {text}")
        else:
            parts.append(text)
    for table in document.tables:
        rendered = _table_markdown([[cell.text for cell in row.cells] for row in table.rows])
        if rendered:
            parts.append(rendered)
    title = document.core_properties.title or None
    return ParsedDocument(pages=[ParsedPage(1, "\n\n".join(parts))], title=title, hidden_text=hidden)


def _parse_pptx(data: bytes) -> ParsedDocument:
    from pptx import Presentation  # noqa: PLC0415

    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"Not a readable PowerPoint file ({type(exc).__name__}).") from exc
    pages: list[ParsedPage] = []
    for number, slide in enumerate(deck.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False) and shape.has_table:
                texts.append(_table_markdown([[c.text for c in row.cells] for row in shape.table.rows]))
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
        if texts:
            pages.append(ParsedPage(number, f"## Slide {number}\n\n" + "\n\n".join(texts)))
    return ParsedDocument(pages=pages)


def _parse_xlsx(data: bytes) -> ParsedDocument:
    from openpyxl import load_workbook  # noqa: PLC0415

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError(f"Not a readable Excel file ({type(exc).__name__}).") from exc
    pages: list[ParsedPage] = []
    for number, sheet in enumerate(workbook.worksheets, start=1):
        rows = [list(row) for _, row in zip(range(_XLSX_MAX_ROWS), sheet.iter_rows(values_only=True), strict=False)]
        rendered = _table_markdown(rows)
        if rendered:
            pages.append(ParsedPage(number, f"## Sheet: {sheet.title}\n\n{rendered}"))
    workbook.close()
    return ParsedDocument(pages=pages)


def _parse_html(data: bytes) -> ParsedDocument:
    from bs4 import BeautifulSoup  # noqa: PLC0415

    soup = BeautifulSoup(_decode_text(data), "html.parser")
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "form", "svg"]):
        tag.decompose()
    hidden = [
        t.get_text(" ", strip=True)[:500]
        for t in soup.select('[style*="display:none"], [style*="display: none"], [hidden]')
    ]
    for tag in soup.select('[style*="display:none"], [style*="display: none"], [hidden]'):
        tag.decompose()
    parts: list[str] = []
    for element in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "table"]):
        if element.name == "table":
            rows = [[c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"])] for tr in element.find_all("tr")]
            parts.append(_table_markdown(rows))
        elif element.name.startswith("h"):
            parts.append("#" * (int(element.name[1]) + 1) + " " + element.get_text(" ", strip=True))
        elif element.name == "li":
            parts.append("- " + element.get_text(" ", strip=True))
        elif element.name == "pre":
            parts.append("```\n" + element.get_text() + "\n```")
        else:
            parts.append(element.get_text(" ", strip=True))
    title = soup.title.get_text(strip=True) if soup.title else None
    return ParsedDocument(
        pages=[ParsedPage(1, "\n\n".join(p for p in parts if p))], title=title, hidden_text=[h for h in hidden if h]
    )


def _parse_image(data: bytes) -> ParsedDocument:
    if not ocr_available():
        raise ParseError("Reading text from images needs the OCR extra (or a vision model). The file was stored.")
    from PIL import Image  # noqa: PLC0415

    text = _ocr_image(Image.open(io.BytesIO(data)))
    if not text.strip():
        raise ParseError("No text was found in this image.")
    return ParsedDocument(pages=[ParsedPage(1, text)], ocr_pages=[1])


def _parse_docling(data: bytes, filename: str) -> ParsedDocument:
    from docling.datamodel.base_models import DocumentStream  # noqa: PLC0415 - optional extra
    from docling.document_converter import DocumentConverter  # noqa: PLC0415

    converter = _ocr_engine_cache.get("docling")
    if converter is None:
        converter = DocumentConverter()
        _ocr_engine_cache["docling"] = converter
    result = converter.convert(DocumentStream(name=filename, stream=io.BytesIO(data)))
    return ParsedDocument(pages=[ParsedPage(1, result.document.export_to_markdown())], meta={"parser": "docling"})


def docling_available() -> bool:
    try:
        import docling  # noqa: F401, PLC0415
    except ImportError:
        return False
    return True


def parse_document(
    data: bytes, mime: str, filename: str, *, backend: str = "auto", max_pages: int = 2000
) -> ParsedDocument:
    use_docling = backend == "docling" or (backend == "auto" and docling_available())
    if use_docling and mime in {_EXTENSIONS["pdf"], _EXTENSIONS["docx"], _EXTENSIONS["pptx"]}:
        if not docling_available():
            raise ParseError("PARSER_BACKEND=docling but Docling is not installed.")
        parsed = _parse_docling(data, filename)
    elif mime == "application/pdf":
        parsed = _parse_pdf(data, max_pages)
    elif mime == _EXTENSIONS["docx"]:
        parsed = _parse_docx(data)
    elif mime == _EXTENSIONS["pptx"]:
        parsed = _parse_pptx(data)
    elif mime == _EXTENSIONS["xlsx"]:
        parsed = _parse_xlsx(data)
    elif mime == "text/html":
        parsed = _parse_html(data)
    elif mime.startswith("image/"):
        parsed = _parse_image(data)
    elif mime.startswith("text/") or mime in {"application/json", "application/yaml"}:
        text = _decode_text(data)
        if mime.startswith("text/x-") or mime in {"application/json", "application/yaml"}:
            language = mime.rsplit("/", 1)[-1].removeprefix("x-")
            text = f"```{language}\n{text}\n```"
        parsed = ParsedDocument(pages=[ParsedPage(1, text)])
    else:
        raise ParseError(
            f"Unsupported file type ({mime}). Supported: PDF, Word, PowerPoint, Excel, HTML, text, code, images."
        )
    parsed.meta.setdefault("parser", "basic")
    # Postgres text cannot hold NUL, and no parser should emit control characters into chunks.
    parsed.pages = [ParsedPage(p.number, _CONTROL_CHARS.sub("", p.markdown)) for p in parsed.pages]
    return parsed
